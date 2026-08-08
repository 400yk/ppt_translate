"""
Tests for where translate_pptx writes its output file.

Regression context: the API and the Celery worker run as separate containers.
They share a volume mounted at UPLOAD_TEMP_DIR (/app/uploads in production),
which is how the worker reads the file the API uploaded. The translated file,
however, was written with a bare tempfile.NamedTemporaryFile() and landed in the
worker's own /tmp.

That made the "local storage" fallback unusable: when both S3 and OSS uploads
failed, the task returned SUCCESS with a path only the worker could see, and
/api/download 404'd from the API container. These tests pin the output into the
shared directory so the fallback can actually serve a file.
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation

from config import UPLOAD_TEMP_DIR
import services.translate_service as translate_service


def _mock_translate(texts, src_lang, dest_lang, batch_size=200, character_batch_size=20000):
    """Stand in for the LLM so these tests make no network calls."""
    return [f'translated {t}' for t in texts], sum(len(t) for t in texts)


def _build_pptx(with_text=True):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if with_text:
        box = slide.shapes.add_textbox(0, 0, 1000000, 500000)
        box.text_frame.text = '日商拓産概要'
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _in_shared_dir(path):
    """The API container can only serve files under the shared volume."""
    shared = os.path.realpath(UPLOAD_TEMP_DIR)
    return os.path.commonpath([os.path.realpath(path), shared]) == shared


def test_translated_file_lands_in_the_shared_upload_dir(monkeypatch=None):
    original = translate_service.gemini_batch_translate_with_size
    translate_service.gemini_batch_translate_with_size = _mock_translate
    try:
        path, _ = translate_service.translate_pptx(_build_pptx(), 'ja', 'en')
    finally:
        translate_service.gemini_batch_translate_with_size = original
    try:
        assert _in_shared_dir(path), f'{path} is not under {UPLOAD_TEMP_DIR}'
        assert os.path.exists(path)
    finally:
        os.path.exists(path) and os.unlink(path)


def test_untranslatable_deck_also_lands_in_the_shared_upload_dir():
    """The early-return path for decks with no translatable text needs it too."""
    path, chars = translate_service.translate_pptx(_build_pptx(with_text=False), 'ja', 'en')
    try:
        assert chars == 0
        assert _in_shared_dir(path), f'{path} is not under {UPLOAD_TEMP_DIR}'
    finally:
        os.path.exists(path) and os.unlink(path)


if __name__ == '__main__':
    failures = 0
    for name, fn in sorted(globals().items()):
        if name.startswith('test_') and callable(fn):
            try:
                fn()
                print(f'PASS {name}')
            except AssertionError as exc:
                failures += 1
                print(f'FAIL {name}: {exc}')
            except Exception as exc:
                failures += 1
                print(f'ERROR {name}: {type(exc).__name__}: {exc}')
    print(f'\n{"FAILED" if failures else "OK"} ({failures} failing)')
    sys.exit(1 if failures else 0)
